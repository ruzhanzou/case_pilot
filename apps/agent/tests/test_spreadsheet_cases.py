from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from openpyxl import Workbook

from casepilot_agent.knowledge import parse_document
from casepilot_agent.reuse import normalize_sources, spreadsheet_import_report
from casepilot_agent.spreadsheet_cases import organize_spreadsheet_cases


def case(**values):
    return {
        "row": 2,
        "title": ["B2"],
        "steps": ["C2", "C3"],
        "expected": ["D3"],
        "original_id": ["A2"],
        "preconditions": [],
        "module": [],
        "priority": [],
        "case_type": [],
        "needs_review": False,
        "reason": "",
        **values,
    }


def provider_for(cases, ignored=None):
    def complete(**kwargs):
        assert kwargs["stage"] == "attachment.cases_organized"
        assert "禁止输出字段正文" in kwargs["instruction"]
        return kwargs["result_type"](cases=cases, ignored=ignored or []), SimpleNamespace()

    return SimpleNamespace(complete=Mock(side_effect=complete))


def workbook(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "登录"
    sheet.append(["标识", "检查事项", "如何操作", "应该看到什么"])
    sheet.append(["0001", "过期验证码拒绝登录", "签发验证码", None])
    sheet.append([None, None, "推进到300秒再登录", "返回E731"])
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:B3")
    path = tmp_path / "自由格式.xlsx"
    book.save(path)
    return parse_document(path, "")


def test_unknown_headers_merged_multiline_cases_keep_exact_cells_and_provenance(tmp_path):
    blocks = workbook(tmp_path)
    original = blocks[0].content
    assert not blocks[0].metadata.get("reference_cases")
    provider = provider_for([case()], [{"rows": [1], "reason": "表头"}])
    organize_spreadsheet_cases(blocks, provider)
    assert blocks[0].content == original
    parsed = blocks[0].metadata["reference_cases"][0]
    assert parsed["original_id"] == "0001"
    assert parsed["title"] == "过期验证码拒绝登录"
    assert parsed["steps"] == [{"action": "签发验证码\n推进到300秒再登录", "expected": "返回E731"}]
    assert [step["expected"] for step in parsed["steps"] if step["expected"]] == ["返回E731"]
    assert parsed["row"] == 2 and parsed["row_end"] == 3
    assert parsed["preconditions"] == ["原文件未提供前置条件"]
    assert parsed["source_cells"]["expected"] == ["D3"]
    evidence = {**blocks[0].metadata, "document_id": "document", "label": "自由格式.xlsx"}
    sources = normalize_sources([], [evidence, evidence])
    assert len(sources) == 1
    assert sources[0]["locator"] == "Sheet 登录，行 2-3；0001"
    assert len(spreadsheet_import_report([evidence, evidence])) == 2


@pytest.mark.parametrize(
    "changes",
    [
        {"expected": []},
        {"expected": ["Z999"]},
        {"row": 3},
        {"needs_review": True, "reason": "预期归属不确定"},
    ],
)
def test_incomplete_ambiguous_and_invented_cells_never_become_reusable(tmp_path, changes):
    blocks = workbook(tmp_path)
    organize_spreadsheet_cases(blocks, provider_for([case(**changes)]))
    assert blocks[0].metadata["reference_cases"] == []
    assert any(not row["valid"] and row["reason"] for row in blocks[0].metadata["case_import_rows"])


def test_failure_retains_standard_extraction_and_visible_warning(tmp_path):
    book = Workbook()
    book.active.append(["用例名称", "测试步骤", "预期结果"])
    book.active.append(["登录成功", "登录", "成功"])
    path = tmp_path / "标准.xlsx"
    book.save(path)
    blocks = parse_document(path, "")
    original = deepcopy(blocks[0].metadata["reference_cases"])
    provider = SimpleNamespace(complete=Mock(side_effect=TimeoutError))
    organize_spreadsheet_cases(blocks, provider)
    assert blocks[0].metadata["reference_cases"] == original
    assert blocks[0].metadata["case_organization"] == "fallback"
    assert any(row.get("method") == "fallback" for row in blocks[0].metadata["case_import_rows"])


def test_invalid_reference_gets_one_bounded_repair_before_validation(tmp_path):
    from casepilot_agent.spreadsheet_cases import SpreadsheetCases

    provider = SimpleNamespace(complete=Mock(side_effect=[
        (SpreadsheetCases(cases=[case(expected=["Z999"])], ignored=[]), None),
        (SpreadsheetCases(cases=[case()], ignored=[]), None),
    ]))
    blocks = workbook(tmp_path)
    organize_spreadsheet_cases(blocks, provider)
    assert len(blocks[0].metadata["reference_cases"]) == 1
    assert provider.complete.call_count == 2
    assert provider.complete.call_args.kwargs["payload"]["invalid_references"] == ["Z999"]


def test_cross_block_case_uses_neighbor_rows_but_is_only_stored_once(tmp_path):
    book = Workbook()
    for row in range(1, 40):
        book.active.append([f"说明{row}"])
    book.active.append(["001", "登录", "打开登录页"])
    book.active.append([None, None, "提交", "进入首页"])
    path = tmp_path / "跨块.xlsx"
    book.save(path)
    blocks = parse_document(path, "")
    provider = provider_for(
        [case(row=40, original_id=["A40"], title=["B40"], steps=["C40", "C41"], expected=["D41"])]
    )
    organize_spreadsheet_cases(blocks, provider)
    cases = [item for block in blocks for item in block.metadata["reference_cases"]]
    assert len(cases) == 1
    assert cases[0]["row_end"] == 41
    assert all(
        "D41" in {key for row in call.kwargs["payload"]["rows"] for key in row["cells"]}
        for call in provider.complete.call_args_list
    )


def test_non_spreadsheet_does_not_call_model(tmp_path):
    path = tmp_path / "需求.txt"
    path.write_text("需求说明")
    blocks = parse_document(path, "")
    provider = provider_for([])
    organize_spreadsheet_cases(blocks, provider)
    provider.complete.assert_not_called()


def test_duplicate_anchors_are_not_silently_overwritten(tmp_path):
    blocks = workbook(tmp_path)
    organize_spreadsheet_cases(blocks, provider_for([case(), case()]))
    assert blocks[0].metadata["reference_cases"] == []


def test_header_anchor_is_derived_from_content_before_deduplication(tmp_path):
    blocks = workbook(tmp_path)
    organize_spreadsheet_cases(blocks, provider_for([case(row=1)]))
    assert blocks[0].metadata["reference_cases"][0]["row"] == 2
    blocks = workbook(tmp_path)
    organize_spreadsheet_cases(blocks, provider_for([case(row=1), case(row=2)]))
    assert blocks[0].metadata["reference_cases"] == []


def test_oversized_context_does_not_truncate_and_claim_success(tmp_path):
    blocks = workbook(tmp_path)
    blocks[0].metadata["spreadsheet_rows"][0]["cells"]["A1"] = "原文" * 30_000
    provider = provider_for([])
    organize_spreadsheet_cases(blocks, provider)
    provider.complete.assert_not_called()
    assert blocks[0].metadata["case_organization"] == "fallback"


def test_unrecognized_excel_stays_at_confirmation_instead_of_silently_generating(monkeypatch):
    from uuid import uuid4

    from test_draft_planning_progress import Store

    from casepilot_agent import tasks
    from casepilot_agent.contracts import UsageMetadata
    from casepilot_agent.reuse import ReuseAnalysis

    store = Store()
    store.job["space_id"] = uuid4()
    store.job["input_payload"].update(generation_strategy="select_and_generate", case_context=[])
    store.search_reuse_cases = Mock(return_value=[])
    store.attach_reuse_collections = lambda connection, items, **kwargs: items

    def persist(connection, job, content):
        store.briefs.append(deepcopy(content))
        return 1

    store.persist_test_brief = persist

    class Provider:
        def complete(self, *, stage, **kwargs):
            assert stage in {"requirement.analyzed", "requirement.extracted"}
            return ReuseAnalysis(
                summary="登录", test_object="账号系统", test_object_specified=True,
            ), UsageMetadata(model="mock")

    evidence = [{"document_id": "document", "sheet": "Sheet1", "label": "cases.xlsx",
                 "case_import_rows": [{"row": 2, "valid": False, "reason": "缺少预期"}]}]
    monkeypatch.setattr(tasks, "JobStore", lambda *args: store)
    monkeypatch.setattr(tasks, "create_provider", lambda *args: Provider())
    monkeypatch.setattr(tasks, "create_embedding_provider", lambda: None)
    monkeypatch.setattr(tasks, "_context_payload", lambda *args, **kwargs: {"evidence": evidence})
    tasks.draft_test_brief(str(store.job["id"]))
    content = store.briefs[0]
    assert content["planning"] == {}
    assert content["asset_selection"]["sources"] == []
    assert content["asset_selection"]["import_report"][0]["reason"] == "缺少预期"


@pytest.mark.parametrize("confirmed", [False, True])
def test_comparison_receives_current_confirmation_separately_from_original_prompt(confirmed):
    from test_compact_generation import request
    from test_reuse import source

    from casepilot_agent.reuse import plan_reuse

    def complete(stage, instruction, payload, result_type):
        assert payload["asset_selection_confirmed"] is confirmed
        assert payload["prompt"] == "请先让我确认，不要自动纳入正式集合"
        assert "不等于自动纳入正式集合" in instruction
        return result_type(decisions=[{
            "point_id": payload["test_points"][0]["id"], "action": "reuse",
            "source_key": payload["sources"][0]["key"], "test_object_match": "same",
            "reason": "内容符合当前规则；正式入库另行确认",
        }])

    result = plan_reuse(request(1).confirmed_plan, "regression", [source()], complete,
                        "请先让我确认，不要自动纳入正式集合", assets_confirmed=confirmed)
    assert result["decisions"][0]["action"] == "reuse"
