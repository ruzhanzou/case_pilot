from copy import deepcopy
from unittest.mock import Mock
from uuid import uuid4

import pytest
from openpyxl import Workbook
from sqlalchemy.dialects import postgresql
from test_compact_generation import execute, request, run

from casepilot_agent.knowledge import parse_document
from casepilot_agent.reuse import (
    ReuseDecisions,
    normalize_sources,
    plan_reuse,
    shortlist,
)
from casepilot_agent.store import JobStore


def quantity_only(stage, *args):
    from casepilot_agent.quantity import OutputQuantity
    assert stage == "output_quantity.interpreted", "Generation must not run for reused assets"
    return OutputQuantity()


def source():
    return {
        "key": "asset:one",
        "kind": "formal",
        "label": "资产 TC-001",
        "locator": "版本 4",
        "revision": "revision-4",
        "snapshot": {
            "title": "有效账号建立会话",
            "module": "登录",
            "case_type": "functional",
            "priority": "P1",
            "preconditions": ["已有有效账号"],
            "steps": [{"action": "输入正确凭据登录", "expected": "建立用户会话"}],
        },
    }


def with_decision(action="reuse", count=2):
    req = request(count)
    decision = {
        "point_id": "P-0",
        "point_snapshot": deepcopy(req.confirmed_plan["test_points"][0]),
        "action": action,
        "source_key": "asset:one",
        "source": source(),
        "reason": "核心登录规则未变",
        "changed_fields": ["actions", "expected"] if action == "modify" else [],
    }
    req.confirmed_plan["reuse"] = {"strategy": "regression", "decisions": [decision]}
    return req


def test_adaptation_prompt_preserves_body_without_copying_system_metadata():
    req = with_decision("modify", 1)
    original = req.confirmed_plan["reuse"]["decisions"][0]["source"]["snapshot"]
    original["steps"][0]["id"] = "source-step-id"
    original["source_refs"] = [{"label": "source identity"}]
    def tracked(stage, instruction, payload, result_type, model_id):
        if stage == "output_quantity.interpreted":
            return execute(stage, instruction, payload, result_type, model_id)
        body = payload["reuse_modifications"][0]["original"]
        assert set(body) == {"title", "preconditions", "steps"}
        assert body["preconditions"] == original["preconditions"]
        assert body["steps"] == [{"action": "输入正确凭据登录", "expected": "建立用户会话"}]
        return execute(stage, instruction, payload, result_type, model_id)
    assert len(run(req, tracked).test_cases) == 1


def test_reuse_preserves_content_and_does_not_call_model_for_reused_slot():
    req, calls = with_decision(), []

    def tracked(*args):
        if args[0] == "output_quantity.interpreted":
            return execute(*args)
        calls.append(args[2]["batch_slots"])
        return execute(*args)

    result = run(req, tracked)
    assert calls == [[2]]
    first = result.test_cases[0]
    assert first.title == source()["snapshot"]["title"]
    assert [s.model_dump() for s in first.steps] == source()["snapshot"]["steps"]
    assert "revision-4" in first.source_refs[-1].excerpt
    assert "复用" in first.source_refs[-1].label


def test_all_reused_never_calls_generation_model():
    result = run(with_decision(count=1), quantity_only)
    assert len(result.test_cases) == 1


def test_modified_case_receives_original_and_retains_source():
    calls = []

    def tracked(*args):
        if args[0] == "output_quantity.interpreted":
            return execute(*args)
        calls.append(args[2]["reuse_modifications"])
        return execute(*args)

    result = run(with_decision("modify", count=1), tracked)
    assert calls[0][0]["original"] == {key: source()["snapshot"][key] for key in ("title", "preconditions", "steps")}
    assert any("修改参考" in ref.label for ref in result.test_cases[0].source_refs)


def test_modified_case_preserves_fields_outside_decision_authorization():
    req = with_decision("modify", count=1)
    decision = req.confirmed_plan["reuse"]["decisions"][0]
    decision["reason"] = "当前规则只要求修改预期结果"
    decision["changed_fields"] = ["expected"]

    def overreaching(stage, instruction, payload, result_type, model_id):
        result = execute(stage, instruction, payload, result_type, model_id)
        if stage == "test_case.generated":
            case = result.test_cases[0]
            case.title = "模型擅自重写的标题"
            case.preconditions = ["模型擅自重写的前置条件"]
            case.steps[0].action = "模型擅自重写的操作"
            case.steps[0].expected = "新的可观察预期"
        return result

    candidate = run(req, overreaching).test_cases[0]
    original = source()["snapshot"]
    assert candidate.title == original["title"]
    assert candidate.preconditions == original["preconditions"]
    assert candidate.steps[0].action == original["steps"][0]["action"]
    assert candidate.steps[0].expected == "新的可观察预期"
    assert candidate.module == original["module"]
    assert candidate.priority == original["priority"]
    assert candidate.case_type == original["case_type"]


def test_changed_point_cannot_use_stale_reuse_decision():
    req = with_decision()
    req.confirmed_plan["test_points"][0]["objective"] = "改成短信登录"
    with pytest.raises(ValueError, match="reuse_plan_changed"):
        run(req)


def test_fresh_strategy_ignores_old_decisions():
    req = with_decision(count=1)
    req.confirmed_plan["reuse"]["strategy"] = "fresh"
    result = run(req)
    assert result.test_cases[0].title != source()["snapshot"]["title"]


def test_one_reused_case_can_cover_two_points_without_duplicate_candidates():
    req = with_decision()
    second = deepcopy(req.confirmed_plan["reuse"]["decisions"][0])
    second.update(point_id="P-1", point_snapshot=deepcopy(req.confirmed_plan["test_points"][1]))
    req.confirmed_plan["reuse"]["decisions"].append(second)
    result = run(req, quantity_only)
    assert len(result.test_cases) == 1
    assert result.test_cases[0].test_point_ids == ["P-0", "P-1"]


@pytest.mark.parametrize(
    "bad", ["unknown_source", "missing_point", "duplicate_point", "new_with_source"]
)
def test_reuse_decisions_reject_invented_sources_and_incomplete_coverage(bad):
    req = request(1)
    row = {"point_id": "P-0", "action": "reuse", "test_object_match": "same", "source_key": "asset:one", "reason": "未变"}
    rows = [row]
    if bad == "unknown_source":
        row["source_key"] = "asset:other-space"
    elif bad == "missing_point":
        rows = []
    elif bad == "duplicate_point":
        rows = [row, row]
    else:
        row["action"] = "new"
    with pytest.raises(ValueError):
        plan_reuse(
            req.confirmed_plan,
            "change",
            [source()],
            lambda *args: ReuseDecisions(decisions=rows),
            "验证码改为2分钟",
        )


def test_planning_priority_cannot_turn_unchanged_asset_into_modify():
    req = request(1)
    result = plan_reuse(
        req.confirmed_plan,
        "change",
        [source()],
        lambda *args: ReuseDecisions(decisions=[{
            "point_id": "P-0",
            "action": "modify",
            "test_object_match": "same",
            "source_key": "asset:one",
            "changed_fields": ["priority"],
            "reason": "规划将该测试点推断为P0，原资产为P1，需要修改优先级",
        }]),
        "登录规则保持不变，复用现有用例",
    )
    decision = result["decisions"][0]
    assert decision["action"] == "reuse"
    assert decision["changed_fields"] == []


def test_explicit_priority_request_authorizes_priority_modification():
    req = request(1)
    result = plan_reuse(
        req.confirmed_plan,
        "change",
        [source()],
        lambda *args: ReuseDecisions(authorized_identity_fields=["priority"], decisions=[{
            "point_id": "P-0",
            "action": "modify",
            "test_object_match": "same",
            "source_key": "asset:one",
            "changed_fields": ["priority"],
            "reason": "用户要求将优先级改为P0",
        }]),
        "把登录用例优先级改为P0",
    )
    decision = result["decisions"][0]
    assert decision["action"] == "modify"
    assert decision["changed_fields"] == ["priority"]


def test_excel_headers_survive_multiple_blocks_and_preserve_identity(tmp_path):
    path = tmp_path / "旧用例.xlsx"
    book = Workbook()
    book.active.title = "登录"
    book.active.append(["用例编号", "用例名称", "前置条件", "测试步骤", "预期结果"])
    for i in range(45):
        book.active.append([f"TC-{i:03}", f"场景{i}", "有账号", "登录", "建立会话"])
    book.create_sheet("需求").append(["编号", "规则"])
    book.save(path)
    blocks = parse_document(path, "application/octet-stream")
    rows = [r for block in blocks for r in block.metadata.get("reference_cases", [])]
    assert len(rows) == 45 and rows[-1]["row"] == 46
    assert rows[0]["original_id"] == "TC-000"
    document_id = str(uuid4())
    evidence = [
        {"document_id": document_id, "chunk_id": str(uuid4()), "label": path.name, **block.metadata}
        for block in blocks
    ]
    # Parent chunk splitting must not duplicate source cases.
    sources = normalize_sources([], evidence + evidence)
    assert len(sources) == 45 and "行 46" in sources[-1]["locator"]


def test_incomplete_excel_case_is_not_fabricated(tmp_path):
    path = tmp_path / "不完整.xlsx"
    book = Workbook()
    book.active.append(["用例名称", "测试步骤", "预期结果"])
    book.active.append(["登录", "登录", ""])
    book.save(path)
    assert (
        normalize_sources(
            [], [{"document_id": "x", "label": "x", **b.metadata} for b in parse_document(path, "")]
        )
        == []
    )


def test_shortlist_has_hard_input_budget_without_truncating_snapshots():
    sources = [{**source(), "key": str(i)} for i in range(100)]
    selected = shortlist(sources, [{"title": "登录"}], limit=10, budget=1000)
    import json

    assert 0 < len(selected) <= 10
    assert sum(len(json.dumps(s, ensure_ascii=False)) for s in selected) <= 1000
    assert all(s["snapshot"] == source()["snapshot"] for s in selected)


def test_asset_search_is_bounded_scoped_and_escapes_wildcards():
    connection = Mock()
    connection.execute.return_value.mappings.return_value = []
    JobStore.__new__(JobStore).search_reuse_cases(connection, space_id=uuid4(), terms=["登录%_"])
    sql = "\n".join(str(call.args[0].compile(dialect=postgresql.dialect()))
                    for call in connection.execute.call_args_list)
    assert "space_id =" in sql and "deleted_at IS NULL" in sql
    assert "current_revision_id" in sql and "LIMIT" in sql and "ESCAPE" in sql


@pytest.mark.parametrize("strategy", ["fresh", "regression", "change"])
def test_draft_workflow_queries_assets_only_for_reuse_and_persists_decisions(monkeypatch, strategy):
    from test_draft_planning_progress import Store
    from test_parallel_planning import features, points

    from casepilot_agent import tasks
    from casepilot_agent.contracts import UsageMetadata
    from casepilot_agent.reuse import ReuseAnalysis

    store = Store()
    store.job["space_id"] = uuid4()
    store.job["input_payload"].update(generation_strategy="generate" if strategy == "fresh" else "select_and_generate", asset_selection_confirmed=strategy != "fresh", confirmed_asset_sources=[source()])
    search = Mock(return_value=[{"ref": "one", "snapshot": source()["snapshot"]}])
    store.search_reuse_cases = search
    store.attach_reuse_collections = lambda connection, items, **kwargs: items

    class Provider:
        def complete(self, *, stage, payload, result_type, **kwargs):
            if stage == "output_quantity.interpreted":
                from casepilot_agent.quantity import OutputQuantity
                return OutputQuantity(test_point_count=12), UsageMetadata(model="fixture")
            if stage in {"requirement.analyzed", "requirement.extracted"}:
                result = ReuseAnalysis(
                    summary="登录",
                    test_object="登录",
                    test_object_specified=True,
                    generation_strategy=strategy,
                    case_search_terms=["登录"],
                )
            elif stage == "feature.generated":
                result = features(6)
            elif stage == "reuse.compared":
                result = result_type(
                    decisions=[
                        {
                            "point_id": p["id"],
                            "action": "reuse", "test_object_match": "same",
                            "source_key": payload["sources"][0]["key"],
                            "reason": "登录规则未变",
                        }
                        for p in payload["test_points"]
                    ]
                )
            else:
                result = points(payload, fill_budget=True)
            return result, UsageMetadata(model="mock")

    monkeypatch.setattr(tasks, "JobStore", lambda *args: store)
    monkeypatch.setattr(tasks, "create_provider", lambda *args: Provider())
    monkeypatch.setattr(tasks, "create_embedding_provider", lambda: None)
    monkeypatch.setattr(tasks, "_context_payload", lambda *args, **kwargs: {"evidence": []})
    tasks.draft_test_brief(str(store.job["id"]))
    reuse = store.briefs[0]["planning"]["reuse"]
    assert reuse["strategy"] == ("fresh" if strategy == "fresh" else "regression")
    assert reuse["summary"] == ({"new": 12} if strategy == "fresh" else {"reuse": 12})
    assert search.call_count == 0
    if strategy != "fresh":
        events = [event["event"] for event in store.events]
        assert events.index("cases.searching") < events.index("reuse.comparing")
        assert events[-1] == "brief.completed"


def test_reference_budget_never_causes_unverified_direct_reuse():
    req = request(1)
    req.confirmed_plan["test_points"][0]["source_refs"] = [{"chunk_id": "chunk", "label": "规则"}]
    req.confirmed_plan["evidence"] = [{"chunk_id": "chunk", "excerpt": "x" * 24001}]
    result = plan_reuse(
        req.confirmed_plan,
        "change",
        [source()],
        lambda *args: ReuseDecisions(
            decisions=[
                {
                    "point_id": "P-0",
                    "action": "reuse", "test_object_match": "same",
                    "source_key": "asset:one",
                    "reason": "未变化",
                }
            ]
        ),
        "新需求",
    )
    assert result["decisions"][0]["action"] == "new"
    assert "source" not in result["decisions"][0]


def test_reuse_comparison_bounds_concurrency_and_batch_payloads():
    from threading import Barrier, Lock

    req = request(36)
    lock, barrier = Lock(), Barrier(3)
    active = peak = calls = 0

    def complete(stage, instruction, payload, result_type):
        nonlocal active, peak, calls
        with lock:
            active += 1
            peak = max(peak, active)
            calls += 1
        assert len(payload["test_points"]) <= 12
        assert len(payload["sources"]) <= 30
        barrier.wait(timeout=5)
        result = result_type(decisions=[{
            "point_id": p["id"], "action": "new", "source_key": None, "reason": "缺口"
        } for p in payload["test_points"]])
        with lock:
            active -= 1
        return result

    result = plan_reuse(req.confirmed_plan, "regression", [source()], complete, "登录回归")
    assert peak == 3 and calls == 3
    assert len(result["decisions"]) == 36


def test_cancelled_reuse_does_not_start_model_work():
    complete = Mock()

    def cancelled():
        raise RuntimeError("cancelled")

    with pytest.raises(RuntimeError, match="cancelled"):
        plan_reuse(request(12).confirmed_plan, "change", [source()], complete, "改版",
                   check_cancelled=cancelled)
    complete.assert_not_called()


def test_excel_preserves_priority_type_and_independent_multiline_checks(tmp_path):
    path = tmp_path / '格式验证.xlsx'
    book = Workbook()
    book.active.append(['用例名称', '测试步骤', '预期结果', '优先级', '用例类型'])
    book.active.append(['校验登录', '输入账号\n提交登录', '出现首页', 'P0', '安全'])
    book.active.append(['缺省字段', '登录', '成功', None, None])
    book.active.append(['无效优先级', '登录', '成功', 'P9', None])
    book.active.append(['缺少预期', '登录', None, 'P1', None])
    book.save(path)
    blocks = parse_document(path, '')
    rows = [row for block in blocks for row in block.metadata.get('reference_cases', [])]
    report = [row for block in blocks for row in block.metadata.get('case_import_rows', [])]
    assert len(rows) == 2
    assert rows[0]['priority'] == 'P0' and rows[0]['case_type'] == '安全'
    assert rows[0]['steps'] == [{'action': '输入账号', 'expected': '出现首页'}, {'action': '提交登录', 'expected': ''}]
    assert rows[1]['default_fields'] == ['priority', 'case_type']
    assert [row['row'] for row in report if not row['valid']] == [4, 5]


def test_import_report_deduplicates_rows_from_split_evidence():
    plan = {'test_points': [], 'evidence': [{'document_id': 'd', 'sheet': 's', 'label': 'old.xlsx',
        'case_import_rows': [{'row': 2, 'valid': False, 'default_fields': [], 'reason': '缺少预期'}]}] * 2}
    result = plan_reuse(plan, 'fresh', [], Mock(), '')
    assert len(result['import_report']) == 1


@pytest.mark.parametrize("match", ["different", "unknown", None])
@pytest.mark.parametrize("action", ["reuse", "modify"])
def test_different_or_unverified_test_object_never_reuses(match, action):
    row = {"point_id": "P-0", "action": action, "source_key": "asset:one", "reason": "similar"}
    if match is not None:
        row["test_object_match"] = match
    result = plan_reuse(request(1).confirmed_plan, "change", [source()],
                        lambda *args: ReuseDecisions(decisions=[row]), "login",
                        test_object="other product")
    decision = result["decisions"][0]
    assert decision["action"] == "new"
    assert decision["source_key"] is None
    assert "source" not in decision


def test_same_test_object_can_be_proposed_for_reuse():
    def complete(stage, instruction, payload, result_type):
        assert payload["test_object"] == "account system"
        return result_type(decisions=[{"point_id": "P-0", "action": "reuse",
            "source_key": "asset:one", "reason": "same product and object",
            "test_object_match": "same"}])
    result = plan_reuse(request(1).confirmed_plan, "change", [source()], complete,
                        "login", test_object="account system")
    assert result["decisions"][0]["action"] == "reuse"


def test_reuse_collection_provenance_is_scoped_and_preserves_case_snapshot():
    asset_id, collection_id = uuid4(), uuid4()
    connection = Mock()
    connection.execute.return_value.mappings.return_value = [
        {"test_case_id": asset_id, "id": collection_id, "name": "Account regression"}
    ]
    items = [{"ref": str(asset_id), "target_type": "formal", "snapshot": source()["snapshot"]}]
    result = JobStore.__new__(JobStore).attach_reuse_collections(
        connection, items, space_id=uuid4()
    )
    assert result[0]["collections"] == [{"id": str(collection_id), "name": "Account regression"}]
    assert "collections" not in items[0]
    assert normalize_sources(result, [])[0]["collections"] == result[0]["collections"]
    sql = str(connection.execute.call_args.args[0].compile(dialect=postgresql.dialect()))
    assert "space_id =" in sql and "deleted_at IS NULL" in sql


def test_excel_reference_preserves_original_identity_and_version():
    import json
    from casepilot_agent.reuse import reuse_ref

    asset = {
        "key": "excel:document:历史:2", "kind": "excel", "label": "历史用例.xlsx",
        "locator": "历史!2", "document_version": 3, "document_checksum": "snapshot-hash",
        "snapshot": {"original_id": "OLD-001"},
    }
    ref = reuse_ref({"source": asset, "action": "modify", "reason": "按当前时长改写"})
    provenance = json.loads(ref.excerpt)
    assert provenance["case_key"] == "OLD-001"
    assert provenance["revision"] == 3
    assert provenance["source_label"] == "历史用例.xlsx"
    assert provenance["source_locator"] == "历史!2"
    assert provenance["document_checksum"] == "snapshot-hash"


def test_legacy_modification_permissions_are_interpreted_by_model():
    from casepilot_agent.compact_generation import ModificationFields
    req = with_decision("modify", count=1)
    req.confirmed_plan["reuse"]["decisions"][0].pop("changed_fields")
    calls = []

    def interpret(stage, instruction, payload, result_type, model_id):
        if stage == "reuse.permissions_interpreted":
            calls.append(payload)
            return ModificationFields(changed_fields=["expected"])
        return execute(stage, instruction, payload, result_type, model_id)

    result = run(req, interpret)
    assert len(calls) == 1
    assert calls[0]["prompt"] == req.prompt
    assert result.test_cases[0].steps[0].action == source()["snapshot"]["steps"][0]["action"]
